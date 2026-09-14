"""tasks.json -> 특정 인수자에게 넘기는 업무만 골라 인수인계서 문서 생성 (docx + html)

사용:  python make_handover.py 홍길동
       python make_handover.py 홍길동 --type in     # 인수(받는) 쪽 문서
"""
import io, json, os, sys
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "인계자료")

# 실명·조직 설정은 names.local.json 에서 읽는다 (공개 저장소 제외 대상)
CFG_PATH = os.path.join(BASE, "names.local.json")
if not os.path.exists(CFG_PATH):
    raise SystemExit("names.local.json 없음 — README 의 '로컬에서 빌드' 참고")
CFG = json.load(io.open(CFG_PATH, encoding="utf-8"))

LINK = CFG["handover_link"]   # 업무 <-> 정기일정 연결 키워드

TN = {"keep": "유지", "out": "인계", "in": "인수", "common": "공통", "tbd": "미정"}


def _decode(src):
    """shots[].src -> docx 삽입용 파일 객체. data URI 이거나 캡처/ 폴더의 상대경로."""
    import base64
    if src.startswith("data:"):
        try:
            return io.BytesIO(base64.b64decode(src.split(",", 1)[1]))
        except Exception:
            return None
    f = os.path.join(BASE, src)
    return f if os.path.exists(f) else None


def linked_events(task, events):
    keys = LINK.get(task["id"], [])
    if not keys:
        return []
    hay = lambda e: (e.get("title", "") + " " + e.get("detail", ""))
    return [e for e in events if any(k in hay(e) for k in keys)]


def build(name, kind="out"):
    db = json.load(io.open(os.path.join(BASE, "tasks.json"), encoding="utf-8"))
    tasks = [t for t in db["tasks"] if t["type"] == kind and name in (t.get("counterpart") or "")]
    if not tasks:
        sys.exit(f"'{name}' 에게 {TN[kind]}하는 업무가 없음")
    events = db.get("events", [])
    giver, taker = (db["meta"]["작성자"], name) if kind == "out" else (name, db["meta"]["작성자"])
    os.makedirs(OUT, exist_ok=True)
    stem = os.path.join(OUT, f"인수인계서_{name}_{date.today().isoformat()}")
    _docx(stem + ".docx", db, tasks, events, giver, taker)
    _html(stem + ".html", db, tasks, events, giver, taker)
    return stem, tasks


def _docx(path, db, tasks, events, giver, taker):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.font.size = Pt(10)
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin = s.right_margin = Cm(2)

    h = d.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("업 무 인 수 인 계 서")
    r.bold = True
    r.font.size = Pt(18)

    sub = d.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(f'{db["meta"]["소속"]} · {db["meta"].get("적용","")}')
    sr.font.size = Pt(9)
    sr.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

    t = d.add_table(rows=2, cols=4)
    t.style = "Table Grid"
    for i, v in enumerate(["인계자", "인수자", "입회자", "일  자"]):
        c = t.rows[0].cells[i]
        c.text = v
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, v in enumerate([giver, taker, "", ""]):
        t.rows[1].cells[i].text = v
        t.rows[1].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph()

    d.add_heading("1. 인계 업무 총괄", level=1)
    head = ["연번", "분야", "업무명", "주기", "기한·마감", "관련 문서 파일명"]
    widths = [Cm(1.1), Cm(1.6), Cm(4.0), Cm(2.0), Cm(3.0), Cm(5.3)]
    s = d.add_table(rows=1, cols=len(head))
    s.style = "Table Grid"
    s.autofit = False
    for i, v in enumerate(head):
        c = s.rows[0].cells[i]
        c.text = v
        c.width = widths[i]
        c.paragraphs[0].runs[0].bold = True
    for i, x in enumerate(tasks, 1):
        row = s.add_row().cells
        vals = [str(i), x.get("cat", ""), x["name"], x.get("cycle", ""),
                x.get("deadline", "") or "-", "\n".join(x.get("files") or []) or ""]
        for j, v in enumerate(vals):
            row[j].text = v
            row[j].width = widths[j]
    fn = d.add_paragraph()
    fr = fn.add_run("※ 관련 문서 파일명이 빈 칸인 업무는 인계 시 사용 서식·대장 파일명을 직접 기재한다.")
    fr.font.size = Pt(8.5)
    fr.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
    d.add_paragraph()

    d.add_heading("2. 업무별 인계 내용", level=1)
    for i, x in enumerate(tasks, 1):
        d.add_heading(f'{i}. {x["name"]}', level=2)

        meta = " | ".join(filter(None, [
            x.get("cat"), x.get("cycle"),
            ("기한 " + x["deadline"]) if x.get("deadline") else "",
            ("시스템 " + ", ".join(x["systems"])) if x.get("systems") else "",
        ]))
        if meta:
            mp = d.add_paragraph()
            mr = mp.add_run(meta)
            mr.font.size = Pt(9)
            mr.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

        def block(title, items, numbered=True, blanks=0):
            p = d.add_paragraph()
            pr = p.add_run(title)
            pr.bold = True
            pr.font.size = Pt(9.5)
            if items:
                for it in items:
                    d.add_paragraph(it, style="List Number" if numbered else "List Bullet")
            else:
                for _ in range(blanks or 1):
                    d.add_paragraph("　", style="List Number" if numbered else "List Bullet")

        d.add_paragraph(x.get("desc") or "　")
        block("처리 절차", x.get("steps") or [], True, 5)
        block("자주 발생하는 이슈 · 주의사항", x.get("traps") or [], False, 2)
        block("관련 문서 파일명 (서식·양식·대장)", x.get("files") or [], False, 3)
        block("접근 권한 · 폴더 위치 · 링크", x.get("assets") or [], False, 2)

        shots = x.get("shots") or []
        if shots:
            sp = d.add_paragraph()
            spr = sp.add_run("업무 캡처 · 화면 사진")
            spr.bold = True
            spr.font.size = Pt(9.5)
            for sh in shots:
                buf = _decode(sh.get("src", ""))
                if not buf:
                    continue
                try:
                    d.add_picture(buf, width=Cm(13))
                except Exception:
                    continue
                if sh.get("t"):
                    cp = d.add_paragraph()
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cr2 = cp.add_run(sh["t"])
                    cr2.font.size = Pt(8.5)
                    cr2.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

        ev = linked_events(x, events)
        if ev:
            block("연동 정기일정", [f'{e["title"]} — {e["due"]} ({e.get("system") or e.get("fund") or "-"})'
                                for e in ev], False)
        if x.get("log"):
            block("작업 기록", [f'[{e["d"]}] {e["t"]}' for e in x["log"]], False)

        cp = d.add_paragraph()
        cr = cp.add_run("인수자 확인 : ________________     인계일 : ______________")
        cr.font.size = Pt(9)
        d.add_paragraph()

    d.add_heading("3. 인계 시 확인 사항", level=1)
    for line in [
        "각 업무의 처리 절차는 인계자가 최초 1회 동행 처리하며 공란을 채운다.",
        "마감일이 있는 업무는 인계 직후 도래하는 첫 마감을 인계자·인수자가 함께 처리한다.",
        "계정·권한이 필요한 업무는 인계일에 접근 권한 이전을 완료한다.",
        "미결 건(처리 중이던 문서·미납·미제출)은 목록으로 별도 첨부한다.",
    ]:
        d.add_paragraph(line, style="List Bullet")
    d.save(path)


def _html(path, db, tasks, events, giver, taker):
    esc = lambda s: (str(s) if s is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    ul = lambda items, cls="": ("<ul class='%s'>" % cls) + "".join(f"<li>{esc(i)}</li>" for i in items) + "</ul>" \
        if items else "<p class='none'>미작성 — 인계 시 함께 작성</p>"

    body = []
    for i, x in enumerate(tasks, 1):
        ev = linked_events(x, events)
        body.append(f"""<section class="item">
<h3>{i}. {esc(x['name'])}</h3>
<p class="meta">{esc(' | '.join(filter(None,[x.get('cat'),x.get('cycle'),
  ('기한 '+x['deadline']) if x.get('deadline') else '',
  ('시스템 '+', '.join(x['systems'])) if x.get('systems') else ''])))}</p>
<h4>업무 개요</h4><p>{esc(x.get('desc') or '—')}</p>
<h4>처리 절차</h4>{('<ol>'+''.join(f'<li>{esc(s)}</li>' for s in x['steps'])+'</ol>') if x.get('steps') else "<p class='none'>미작성 — 인계 시 함께 작성</p>"}
<h4>자주 발생하는 이슈 · 주의사항</h4>{ul(x.get('traps') or [], 'trap')}
<h4>관련 문서 파일명 (서식·양식·대장)</h4>{('<ul>'+''.join(f'<li>{esc(f)}</li>' for f in x['files'])+'</ul>') if x.get('files') else "<p class='blankline'>&nbsp;</p><p class='blankline'>&nbsp;</p>"}
<h4>접근 권한 · 폴더 위치 · 링크</h4>{ul(x.get('assets') or [])}
{('<h4>업무 캡처 · 화면 사진</h4><div class="shots">' + ''.join(f'<figure><img src="{esc(sh.get("src",""))}" alt=""><figcaption>{esc(sh.get("t",""))}</figcaption></figure>' for sh in x['shots']) + '</div>') if x.get('shots') else ''}
{('<h4>연동 정기일정</h4>' + ul([f"{e['title']} — {e['due']}" for e in ev])) if ev else ''}
{('<h4>작업 기록</h4>' + ul([f"[{e['d']}] {e['t']}" for e in x['log']])) if x.get('log') else ''}
<p class="sign">인수자 확인 <span></span> 인계일 <span></span></p>
</section>""")

    rows = "".join(f"<tr><td>{i}</td><td>{esc(x.get('cat'))}</td><td>{esc(x['name'])}</td>"
                   f"<td>{esc(x.get('cycle'))}</td><td>{esc(x.get('deadline') or '-')}</td>"
                   f"<td class='fcell'>{'<br>'.join(esc(f) for f in (x.get('files') or []))}</td></tr>"
                   for i, x in enumerate(tasks, 1))

    io.open(path, "w", encoding="utf-8").write(f"""<!doctype html><html lang="ko"><meta charset="utf-8">
<title>인수인계서 · {esc(taker)}</title>
<style>
body{{font:11pt/1.7 "맑은 고딕",system-ui,sans-serif;color:#111;max-width:800px;margin:0 auto;padding:28px 20px;background:#fff}}
h1{{text-align:center;letter-spacing:.25em;font-size:21px;margin:0 0 4px}}
.top{{text-align:center;color:#777;font-size:12px;margin-bottom:20px}}
table{{border-collapse:collapse;width:100%;font-size:12px;margin-bottom:24px}}
th,td{{border:1px solid #999;padding:6px 8px;text-align:left}}
th{{background:#f2f0ec}}
h2{{font-size:15px;border-bottom:2px solid #111;padding-bottom:5px;margin:26px 0 12px}}
.item{{border:1px solid #ccc;border-radius:6px;padding:14px 16px;margin-bottom:14px;break-inside:avoid}}
.item h3{{margin:0 0 4px;font-size:14.5px}}
.meta{{color:#777;font-size:11px;margin:0 0 10px}}
h4{{font-size:11px;color:#555;letter-spacing:.05em;margin:12px 0 4px}}
ol,ul{{margin:0;padding-left:20px;font-size:12.5px}}
ul.trap li{{color:#b91c1c}}
.none{{color:#999;font-style:italic;font-size:12px;margin:0}}
.fcell{{font-size:11px;color:#333}}
.foot{{font-size:10.5px;color:#777;margin:-16px 0 22px}}
.blankline{{border-bottom:1px dotted #bbb;margin:0 0 6px;height:17px}}
.shots{{display:flex;flex-wrap:wrap;gap:10px;margin-top:4px}}
.shots figure{{margin:0;max-width:340px}}
.shots img{{max-width:100%;border:1px solid #ccc;border-radius:5px;display:block}}
.shots figcaption{{font-size:10.5px;color:#777;margin-top:3px}}
.sign{{margin:12px 0 0;font-size:11px;color:#555}}
.sign span{{display:inline-block;width:120px;border-bottom:1px solid #999;margin:0 14px 0 4px}}
@media print{{body{{padding:0}} .item{{border-color:#999}}}}
</style>
<h1>업 무 인 수 인 계 서</h1>
<div class="top">{esc(db['meta']['소속'])} · {esc(db['meta'].get('적용',''))}</div>
<table><tr><th>인계자</th><th>인수자</th><th>입회자</th><th>일자</th></tr>
<tr><td>{esc(giver)}</td><td>{esc(taker)}</td><td></td><td></td></tr></table>
<h2>1. 인계 업무 총괄 ({len(tasks)}건)</h2>
<table><tr><th style="width:34px">연번</th><th style="width:54px">분야</th><th>업무명</th>
<th style="width:88px">주기</th><th style="width:130px">기한·마감</th><th style="width:180px">관련 문서 파일명</th></tr>{rows}</table>
<p class="foot">※ 관련 문서 파일명이 빈 칸인 업무는 인계 시 사용 서식·대장 파일명을 직접 기재한다.</p>
<h2>2. 업무별 인계 내용</h2>
{''.join(body)}
<h2>3. 인계 시 확인 사항</h2>
<ul>
<li>각 업무의 처리 절차는 인계자가 최초 1회 동행 처리하며 공란을 채운다.</li>
<li>마감일이 있는 업무는 인계 직후 도래하는 첫 마감을 인계자·인수자가 함께 처리한다.</li>
<li>계정·권한이 필요한 업무는 인계일에 접근 권한 이전을 완료한다.</li>
<li>미결 건(처리 중이던 문서·미납·미제출)은 목록으로 별도 첨부한다.</li>
</ul>
</html>""")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("사용: python make_handover.py <이름> [--type in]")
    who = sys.argv[1]
    kind = sys.argv[sys.argv.index("--type") + 1] if "--type" in sys.argv else "out"
    stem, tasks = build(who, kind)
    print(f"{who} — {TN[kind]} 업무 {len(tasks)}건")
    for t in tasks:
        n = len(t.get("steps") or [])
        print(f"  · {t['name']}  (절차 {n}단계{'' if n else ' — 공란'})")
    print("생성:", stem + ".docx")
    print("     ", stem + ".html")
